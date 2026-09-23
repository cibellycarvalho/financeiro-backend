import io
import zipfile
from unittest.mock import patch

import openpyxl
import requests

import pacote_zip
import storage

COMPRA = {
    "fornecedor": "FLAVIA",
    "data_compra": "2026-08-17",
    "numero_pedido": "1234",
    "valor": 30000.0,
    "nf_path": "forn/notas/ped-1.pdf",
    "pedido_path": "forn/pedidos/ped-1.pdf",
    "comprovantes": [
        {"path": "forn/pagamentos/pg-1.pdf", "data": "2026-09-22", "valor": 30000.0},
    ],
}


def _nomes(zbytes):
    return zipfile.ZipFile(io.BytesIO(zbytes)).namelist()


def test_pasta_e_da_compra_mesmo_com_pix_do_mes_seguinte():
    with patch("storage.baixar", return_value=b"%PDF"):
        nomes = _nomes(pacote_zip.montar([COMPRA]))
    assert "FLAVIA/17-08 pedido 1234/nota-fiscal.pdf" in nomes
    assert "FLAVIA/17-08 pedido 1234/pedido.pdf" in nomes
    assert "FLAVIA/17-08 pedido 1234/comprovante 22-09 R$ 30.000,00.pdf" in nomes


def test_mesmo_comprovante_aparece_nas_duas_compras():
    a = dict(COMPRA, numero_pedido="10", valor=30000.0)
    b = dict(COMPRA, numero_pedido="11", valor=19310.0, nf_path=None, pedido_path=None,
             comprovantes=[{"path": "forn/pagamentos/pg-1.pdf", "data": "2026-09-22", "valor": 19310.0}])
    with patch("storage.baixar", return_value=b"%PDF"):
        nomes = _nomes(pacote_zip.montar([a, b]))
    assert "FLAVIA/17-08 pedido 10/comprovante 22-09 R$ 30.000,00.pdf" in nomes
    assert "FLAVIA/17-08 pedido 11/comprovante 22-09 R$ 19.310,00.pdf" in nomes


def test_arquivo_que_falha_vira_bilhete_em_vez_de_derrubar_o_zip():
    import storage as st
    with patch("storage.baixar", side_effect=st.StorageErro("sumiu")):
        conteudo = pacote_zip.montar([COMPRA])
    nomes = _nomes(conteudo)
    assert any(n.startswith("FLAVIA/17-08 pedido 1234/FALTOU") for n in nomes)


def test_compra_sem_arquivo_nenhum_ainda_tem_pasta_e_linha_no_resumo():
    vazia = dict(COMPRA, nf_path=None, pedido_path=None, comprovantes=[])
    conteudo = pacote_zip.montar([vazia])
    nomes = _nomes(conteudo)
    assert any(n.startswith("FLAVIA/17-08 pedido 1234/") for n in nomes)
    assert "resumo.xlsx" in nomes


def test_resumo_marca_o_que_falta():
    vazia = dict(COMPRA, nf_path=None, comprovantes=[])
    conteudo = pacote_zip.montar([vazia])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        planilha = openpyxl.load_workbook(io.BytesIO(z.read("resumo.xlsx")))
    linhas = list(planilha.active.values)
    assert linhas[0][0] == "Fornecedor"
    assert "não" in linhas[1]    # tem nota = não


def test_cabecalho_tem_data_compra_ou_vencimento_e_origem():
    conteudo = pacote_zip.montar([COMPRA])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        planilha = openpyxl.load_workbook(io.BytesIO(z.read("resumo.xlsx")))
    cabecalho = list(planilha.active.values)[0]
    assert cabecalho == (
        "Fornecedor", "Data (compra ou vencimento)", "Nº do pedido", "Valor", "Pago",
        "Pagamentos", "Em aberto", "Tem nota", "Tem comprovante", "Origem",
    )


# --- Sanitização de nome de pasta/arquivo (Fix round 1, item 1) ------------

def test_numero_de_pedido_com_barra_nao_cria_pasta_extra():
    # Formato real da Flávia: "2026/9001" — a "/" não pode virar nível de pasta.
    com_barra = dict(COMPRA, numero_pedido="2026/9001", comprovantes=[])
    nomes = _nomes(pacote_zip.montar([com_barra]))
    assert any(n.startswith("FLAVIA/17-08 pedido 2026-9001/") for n in nomes)
    assert not any("2026/9001" in n for n in nomes)


def test_descricao_com_dois_pontos_nao_gera_caractere_ilegal_no_windows():
    com_dois_pontos = dict(COMPRA, numero_pedido="NF: 123", comprovantes=[])
    nomes = _nomes(pacote_zip.montar([com_dois_pontos]))
    assert any(n.startswith("FLAVIA/17-08 pedido NF- 123/") for n in nomes)
    assert not any(":" in n for n in nomes)


def test_numero_de_pedido_ponto_ponto_nao_faz_zip_slip():
    malicioso = dict(COMPRA, numero_pedido="..", comprovantes=[])
    nomes = _nomes(pacote_zip.montar([malicioso]))
    assert not any("/../" in n or n.startswith("../") for n in nomes)
    assert any(n.startswith("FLAVIA/17-08 pedido -/") for n in nomes)


def test_sanitizar_corta_em_60_caracteres_e_colapsa_espaços():
    texto = "a" * 100
    assert len(pacote_zip._sanitizar(texto)) == 60
    assert pacote_zip._sanitizar("a    b") == "a b"
    assert pacote_zip._sanitizar(".") == "-"
    assert pacote_zip._sanitizar("..") == "-"


# --- Desempate de pasta duplicada (Fix round 1, item 2) --------------------

def test_duas_compras_mesmo_fornecedor_mesmo_dia_sem_numero_sobrevivem():
    a = dict(COMPRA, id="aaaaaa11", numero_pedido=None, comprovantes=[])
    b = dict(COMPRA, id="bbbbbb22", numero_pedido=None, comprovantes=[])
    with patch("storage.baixar", return_value=b"%PDF"):
        nomes = _nomes(pacote_zip.montar([a, b]))
    assert "FLAVIA/17-08 pedido sem número/nota-fiscal.pdf" in nomes
    assert "FLAVIA/17-08 pedido sem número (bbbbbb)/nota-fiscal.pdf" in nomes


def test_duas_compras_duplicadas_sem_id_ainda_assim_nao_se_pisam():
    a = dict(COMPRA, numero_pedido=None, comprovantes=[])
    b = dict(COMPRA, numero_pedido=None, comprovantes=[])
    conteudo = pacote_zip.montar([a, b])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        pastas = {n.rsplit("/", 1)[0] for n in z.namelist() if n != "resumo.xlsx"}
    assert len(pastas) == 2


def test_montar_e_deterministico_mesma_lista_da_as_mesmas_pastas():
    # A mesma lista, montada duas vezes, tem que sair idêntica — o desempate
    # não pode depender de nada além da ordem da lista recebida.
    a = dict(COMPRA, id="aaaaaa11", numero_pedido=None, comprovantes=[])
    b = dict(COMPRA, id="bbbbbb22", numero_pedido=None, comprovantes=[])
    compras = [a, b]
    pastas1 = sorted(n.rsplit("/", 1)[0] for n in _nomes(pacote_zip.montar(compras)) if n != "resumo.xlsx")
    pastas2 = sorted(n.rsplit("/", 1)[0] for n in _nomes(pacote_zip.montar(compras)) if n != "resumo.xlsx")
    assert pastas1 == pastas2


def test_quem_fica_sem_sufixo_depende_da_ordem_da_lista():
    # Por isso a query em routes/pacote.py precisa de ORDER BY estável (com
    # id no fim): montar() em si é determinístico, mas decide "quem é o
    # primeiro" pela ordem em que a lista chega — se o Postgres mandar a
    # lista em ordem diferente a cada vez, a pasta sem sufixo troca de dona.
    a = dict(COMPRA, id="aaaaaa11", numero_pedido=None, comprovantes=[])
    b = dict(COMPRA, id="bbbbbb22", numero_pedido=None, comprovantes=[])
    with patch("storage.baixar", return_value=b"%PDF"):
        nomes_ab = _nomes(pacote_zip.montar([a, b]))
        nomes_ba = _nomes(pacote_zip.montar([b, a]))
    assert "FLAVIA/17-08 pedido sem número/nota-fiscal.pdf" in nomes_ab
    assert "FLAVIA/17-08 pedido sem número (bbbbbb)/nota-fiscal.pdf" in nomes_ab
    assert "FLAVIA/17-08 pedido sem número/nota-fiscal.pdf" in nomes_ba
    assert "FLAVIA/17-08 pedido sem número (aaaaaa)/nota-fiscal.pdf" in nomes_ba


def test_queries_do_mes_desempatam_por_id_no_order_by(mocker):
    # Garante a correção real: sem "id" no fim do ORDER BY, o Postgres não
    # garante a mesma ordem entre execuções, e o teste acima mostra o efeito
    # disso — a pasta sem sufixo trocaria de dona a cada download do mesmo mês.
    query_mock = mocker.patch("routes.pacote.db.query", return_value=[])
    import routes.pacote as rp
    rp._compras_do_mes("2026-08-01", "2026-08-31")
    sql_pedidos = query_mock.call_args_list[0].args[0]
    sql_contas = query_mock.call_args_list[1].args[0]
    assert "ORDER BY f.nome, p.data_pedido, p.id" in sql_pedidos
    assert "ORDER BY c.vencimento, c.id" in sql_contas


# --- Falha de rede não derruba o zip (Fix round 1, item 3) -----------------

def test_falha_de_rede_no_download_vira_bilhete_faltou():
    with patch("storage.requests.get", side_effect=requests.exceptions.ReadTimeout("timeout")):
        conteudo = pacote_zip.montar([COMPRA])
    nomes = _nomes(conteudo)
    assert any(n.startswith("FLAVIA/17-08 pedido 1234/FALTOU") for n in nomes)


# --- Compra paga sem amarração de Pix (Fix round 1, item 4/5) --------------

def test_pedido_pago_sem_amarracao_aparece_pago_no_resumo():
    paga_sem_pix = dict(COMPRA, comprovantes=[], pago_em="2026-08-20")
    conteudo = pacote_zip.montar([paga_sem_pix])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        planilha = openpyxl.load_workbook(io.BytesIO(z.read("resumo.xlsx")))
    linhas = list(planilha.active.values)
    linha = linhas[1]
    # Pago == Valor, Em aberto == 0, mas não existe arquivo de comprovante.
    assert linha[4] == linha[3] == 30000.0
    assert linha[6] == 0.0
    assert linha[8] == "não"


def test_conta_a_pagar_paga_sem_comprovante_aparece_paga_no_resumo():
    conta = dict(COMPRA, comprovantes=[], pago=True, origem="conta a pagar")
    conteudo = pacote_zip.montar([conta])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        planilha = openpyxl.load_workbook(io.BytesIO(z.read("resumo.xlsx")))
    linha = list(planilha.active.values)[1]
    assert linha[4] == linha[3] == 30000.0
    assert linha[6] == 0.0
    assert linha[8] == "não"
    assert linha[9] == "conta a pagar"


def test_conta_a_pagar_nao_paga_continua_em_aberto():
    conta = dict(COMPRA, comprovantes=[], pago=False, origem="conta a pagar")
    conteudo = pacote_zip.montar([conta])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        planilha = openpyxl.load_workbook(io.BytesIO(z.read("resumo.xlsx")))
    linha = list(planilha.active.values)[1]
    assert linha[4] == 0.0
    assert linha[6] == 30000.0


# --- Extensão desconhecida e arredondamento (Fix round 1, item 9/10) -------

def test_extensao_fora_da_lista_vira_bin():
    estranha = dict(COMPRA, nf_path="forn/notas/nota.exe", pedido_path=None, comprovantes=[])
    with patch("storage.baixar", return_value=b"x"):
        nomes = _nomes(pacote_zip.montar([estranha]))
    assert "FLAVIA/17-08 pedido 1234/nota-fiscal.bin" in nomes


def test_path_sem_ponto_vira_bin():
    sem_ext = dict(COMPRA, nf_path="forn/notas/notasemextensao", pedido_path=None, comprovantes=[])
    with patch("storage.baixar", return_value=b"x"):
        nomes = _nomes(pacote_zip.montar([sem_ext]))
    assert "FLAVIA/17-08 pedido 1234/nota-fiscal.bin" in nomes


def test_em_aberto_arredonda_residuo_de_ponto_flutuante():
    imprecisa = dict(
        COMPRA, valor=0.3, comprovantes=[
            {"path": None, "data": "2026-09-22", "valor": 0.1},
            {"path": None, "data": "2026-09-22", "valor": 0.1},
            {"path": None, "data": "2026-09-22", "valor": 0.1},
        ],
    )
    conteudo = pacote_zip.montar([imprecisa])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        planilha = openpyxl.load_workbook(io.BytesIO(z.read("resumo.xlsx")))
    linha = list(planilha.active.values)[1]
    assert linha[6] == 0.0


# --- Pasta vazia grava o bit de diretório (Fix round 1, item 11) -----------

def test_pasta_vazia_e_marcada_como_diretorio_de_verdade():
    vazia = dict(COMPRA, nf_path=None, pedido_path=None, comprovantes=[])
    conteudo = pacote_zip.montar([vazia])
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        info = z.getinfo("FLAVIA/17-08 pedido 1234/")
    assert info.external_attr == 0o40755 << 16


def test_rota_devolve_zip(client, admin_headers):
    linhas = [{
        "fornecedor": "FLAVIA", "data_compra": "2026-08-17", "numero_pedido": "1234",
        "valor": 30000.0, "nf_path": None, "pedido_path": None, "comprovantes": [],
    }]
    with patch("routes.pacote._compras_do_mes", return_value=linhas):
        resp = client.get("/api/pacote/compras/2026-08", headers=admin_headers)
    assert resp.status_code == 200
    assert resp.mimetype == "application/zip"
    assert "compras-2026-08.zip" in resp.headers["Content-Disposition"]


def test_mes_sem_compra_avisa_em_vez_de_zip_vazio(client, admin_headers):
    with patch("routes.pacote._compras_do_mes", return_value=[]):
        resp = client.get("/api/pacote/compras/2026-08", headers=admin_headers)
    assert resp.status_code == 404
    assert "Nenhuma compra" in resp.get_json()["error"]


def test_mes_em_formato_errado_da_400(client, admin_headers):
    resp = client.get("/api/pacote/compras/agosto", headers=admin_headers)
    assert resp.status_code == 400


def test_ano_fora_do_intervalo_da_400(client, admin_headers):
    resp = client.get("/api/pacote/compras/0000-01", headers=admin_headers)
    assert resp.status_code == 400
